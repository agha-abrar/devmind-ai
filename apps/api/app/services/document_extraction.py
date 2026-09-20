"""
Text extraction for uploaded documents. Each extractor takes raw bytes and
returns plain text ready for the document chunker.
"""
from __future__ import annotations

import io

import docx
from pypdf import PdfReader

SUPPORTED_DOCUMENT_TYPES = {"pdf", "txt", "md", "markdown", "docx"}


class DocumentExtractionError(Exception):
    pass


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise DocumentExtractionError(f"Could not read PDF: {exc}") from exc

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise DocumentExtractionError("PDF is password-protected.") from exc

    pages_text = []
    for page in reader.pages:
        try:
            pages_text.append(page.extract_text() or "")
        except Exception:
            continue

    text = "\n\n".join(pages_text).strip()
    if not text:
        raise DocumentExtractionError(
            "No extractable text found (the PDF may be scanned images without OCR)."
        )
    return text


def _extract_docx(data: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentExtractionError(f"Could not read DOCX: {exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)

    text = "\n\n".join(parts).strip()
    if not text:
        raise DocumentExtractionError("No extractable text found in DOCX file.")
    return text


def _extract_plain_text(data: bytes) -> str:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = data.decode("latin-1")
        except Exception as exc:
            raise DocumentExtractionError("Could not decode text file.") from exc

    if not text.strip():
        raise DocumentExtractionError("File is empty.")
    return text


def extract_text(data: bytes, file_type: str) -> str:
    file_type = file_type.lower().lstrip(".")
    if file_type == "pdf":
        return _extract_pdf(data)
    if file_type == "docx":
        return _extract_docx(data)
    if file_type in ("txt", "md", "markdown"):
        return _extract_plain_text(data)
    raise DocumentExtractionError(f"Unsupported document type: .{file_type}")
